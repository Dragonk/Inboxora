#!/usr/bin/env python3
"""Regenerate synthetic attachment regression fixtures; not needed by CI.
Requires reportlab, Pillow, pypdf and python-docx in an isolated Python environment.
All content is synthetic. Upstream Office encryption fixtures are retained separately
with their MIT license and pinned provenance.json.
"""
from pathlib import Path
from io import BytesIO
import email.policy
from email.message import EmailMessage
import json, math, re, struct, subprocess, wave, zipfile
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.lib.pdfencrypt import StandardEncryption
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, NumberObject
from docx import Document
from docx.shared import Inches

root = Path(__file__).resolve().parents[1] / 'backend/fixtures/attachments'
root.mkdir(parents=True, exist_ok=True)
image = Image.new('RGB', (180, 100), '#e8edf3')
draw = ImageDraw.Draw(image); draw.rectangle((15,15,165,85), outline='#305c7c', width=3); draw.text((30,40), 'Inboxora fixture', fill='#18334b')
for suffix, format in [('png','PNG'),('jpg','JPEG'),('gif','GIF'),('webp','WEBP'),('avif','AVIF'),('bmp','BMP'),('tiff','TIFF')]:
    image.save(root/f'image.{suffix}', format=format)
second = Image.new('RGB', (100, 180), '#daeedd'); ImageDraw.Draw(second).text((15,80), 'Second image', fill='#1a4930'); second.save(root/'second.png')
(root/'image.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="180" height="100"><rect width="180" height="100" fill="#e8edf3"/><text x="12" y="55">Inboxora SVG</text></svg>')
(root/'hostile.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="180" height="100" onload="window.__attachmentExecuted=true"><script>window.__attachmentExecuted=true</script><foreignObject><iframe src="https://attachment-tracker.example.test/iframe"/></foreignObject><image href="https://attachment-tracker.example.test/image.png"/><rect width="180" height="100" fill="url(https://attachment-tracker.example.test/fill)"/><text x="10" y="50">Safe SVG text</text></svg>')

pdf = canvas.Canvas(str(root/'hundred-pages.pdf'), pagesize=(612,792), invariant=True)
pdf.setTitle('Inboxora 100-page preview fixture')
for page in range(1,101):
    pdf.bookmarkPage(f'page-{page}')
    if page == 1: pdf.addOutlineEntry('Inboxora document', 'page-1', level=0)
    if page == 87: pdf.addOutlineEntry('Jump to page 87', 'page-87', level=1)
    if page == 88: pdf.addOutlineEntry('Following section', 'page-88', level=1)
    pdf.setFont('Helvetica', 22); pdf.drawString(50,735,f'Inboxora preview page {page}')
    pdf.setFont('Helvetica', 12); pdf.drawString(50,700,f'Unique search token invoice-{page:03d}')
    pdf.drawString(50,675,'Selectable PDF text. This file contains no scripts or external links.')
    pdf.rect(50,560,512,80); pdf.drawImage(str(root/'image.png'),50,410,width=180,height=100)
    pdf.showPage()
pdf.save()
secure = canvas.Canvas(str(root/'password.pdf'), pagesize=(612,792), invariant=True, encrypt=StandardEncryption('preview-password', ownerPassword='fixture-owner-password', strength=128))
secure.drawString(50,730,'Unlocked PDF fixture'); secure.showPage(); secure.save()
writer=PdfWriter(); writer.add_page(PdfReader(str(root/'hundred-pages.pdf')).pages[0])
field=DictionaryObject({NameObject('/Type'):NameObject('/Annot'),NameObject('/Subtype'):NameObject('/Widget'),NameObject('/FT'):NameObject('/Sig'),NameObject('/T'):TextStringObject('EmptySignature'),NameObject('/Rect'):ArrayObject([NumberObject(x) for x in [50,200,250,250]])})
field_ref=writer._add_object(field); writer.pages[0][NameObject('/Annots')]=ArrayObject([field_ref]); writer._root_object[NameObject('/AcroForm')]=DictionaryObject({NameObject('/Fields'):ArrayObject([field_ref]),NameObject('/SigFlags'):NumberObject(3)})
with (root/'empty-signature.pdf').open('wb') as output: writer.write(output)

word=Document(); word.add_heading('Inboxora DOCX preview',0); word.add_paragraph('Formatted document fixture with an embedded image and a table.')
word.add_picture(str(root/'image.png'),width=Inches(2))
table=word.add_table(rows=1, cols=2);table.style='Table Grid';table.rows[0].cells[0].text='Item';table.rows[0].cells[1].text='Value'
cells=table.add_row().cells;cells[0].text='Fixture';cells[1].text='123.45'
word.save(root/'document.docx')
# A related document includes external relationships and an altChunk. Those must
# be removed before the renderer constructs any DOM element.
with zipfile.ZipFile(root/'document.docx') as source, zipfile.ZipFile(root/'external-document.docx','w',zipfile.ZIP_DEFLATED) as target:
    for entry in source.infolist():
        data=source.read(entry.filename)
        if entry.filename=='word/_rels/document.xml.rels':
            data=data.replace(b'</Relationships>',b'<Relationship Id="rIdExternal" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" TargetMode="External" Target="https://attachment-tracker.example.test/docx.png"/><Relationship Id="rIdChunk" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk" TargetMode="External" Target="https://attachment-tracker.example.test/chunk.html"/></Relationships>')
        if entry.filename=='word/document.xml':
            data, count = re.subn(rb'(<a:blip\b[^>]*\s)r:embed="[^"]+"', rb'\1r:link="rIdExternal"', data, count=1)
            assert count == 1, 'External image fixture must reference its relationship'
            data=data.replace(b'</w:body>',b'<w:altChunk r:id="rIdChunk"/></w:body>')
        target.writestr(entry,data)

texts={
 'notes.md':'# Inboxora Markdown\n\nReadable notes.\n\n```mermaid\nflowchart LR\n  Inbox --> Preview\n```\n',
 'invalid-mermaid.md':'```mermaid\n%%{init: {"securityLevel": "loose"}}%%\nflowchart LR\nA-->B\n```\n',
 'data.jsonc':'{\n // Preserve this comment\n "large": 9007199254740993, "name": "Zażółć gęślą jaźń",\n}\n',
 'data.json':'{"name":"Inboxora","items":[1,2,3]}',
 'invalid.json':'{"broken": [1,2,}',
 'data.xml':'<root><item id="1">Inboxora XML</item><item id="2">Second</item></root>',
 'invalid.xml':'<root><broken></root>',
 'data.csv':'Name,Amount,Note\n"Fixture, One",1234.50,"quoted, comma"\nSecond,12.00,"line one\nline two"\n',
 'data.tsv':'Name\tAmount\nFixture\t42\n',
 'text.txt':'Inboxora text fixture\nneedle first\nsecond needle\nZażółć gęślą jaźń\n',
 'document.html':'<h1>Safe HTML fixture</h1><script>window.__attachmentExecuted=true</script><img src="https://attachment-tracker.example.test/html.png"><p style="background-image:url(https://attachment-tracker.example.test/background.png)">No remote images</p><a href="javascript:alert(1)">Not executable</a>',
 'events.ics':'BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//Inboxora//Preview fixture//EN\r\nBEGIN:VEVENT\r\nUID:preview-event-1\r\nDTSTART:20261001T100000Z\r\nDTEND:20261001T110000Z\r\nSUMMARY:First fixture event\r\nEND:VEVENT\r\nBEGIN:VEVENT\r\nUID:preview-event-2\r\nDTSTART:20261002T100000Z\r\nDTEND:20261002T110000Z\r\nSUMMARY:Second fixture event\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n',
 'contacts.vcf':'BEGIN:VCARD\r\nVERSION:3.0\r\nUID:preview-contact-1\r\nFN:First Fixture\r\nEMAIL:first@example.test\r\nEND:VCARD\r\nBEGIN:VCARD\r\nVERSION:3.0\r\nUID:preview-contact-2\r\nFN:Second Fixture\r\nEMAIL:second@example.test\r\nEND:VCARD\r\n',
 'script.sh':'#!/bin/sh\n# Passive attachment fixture, never executed.\nprintf "fixture\\n"\n',
}
for name,text in texts.items(): (root/name).write_text(text,encoding='utf-8')
(root/'polish.txt').write_bytes('Zażółć gęślą jaźń'.encode('cp1250'))
for name in ['presentation.pptx','legacy.doc','legacy.ppt','document.odt','presentation.odp','unknown.bin']:
    (root/name).write_bytes(b'Download-only attachment fixture\x00')

message=EmailMessage(policy=email.policy.SMTP);message['From']='sender@example.test';message['To']='recipient@example.test';message['Subject']='EML fixture';message.set_content('EML plain text body')
message.add_alternative('<h1>EML safe body</h1><img src="https://attachment-tracker.example.test/eml.png"><script>window.__attachmentExecuted=true</script>',subtype='html')
message.add_attachment(b'Nested EML attachment bytes\n',maintype='text',subtype='plain',filename='inner.txt')
message.add_attachment(texts['script.sh'].encode(),maintype='application',subtype='x-sh',filename='inner.sh')
(root/'message.eml').write_bytes(message.as_bytes())
with zipfile.ZipFile(root/'archive.zip','w',zipfile.ZIP_DEFLATED) as archive:
    for name in ['hundred-pages.pdf','image.png','notes.md','data.jsonc','document.docx','script.sh']: archive.write(root/name,name)
with zipfile.ZipFile(root/'nested.zip','w',zipfile.ZIP_DEFLATED) as archive: archive.write(root/'archive.zip','inner/archive.zip')
with zipfile.ZipFile(root/'unsafe-path.zip','w',zipfile.ZIP_DEFLATED) as archive: archive.writestr('../escape.txt','Do not extract this path')
with zipfile.ZipFile(root/'too-many.zip','w',zipfile.ZIP_DEFLATED) as archive:
    for i in range(501): archive.writestr(f'entry-{i}.txt','small')
with zipfile.ZipFile(root/'symlink.zip','w') as archive:
    entry=zipfile.ZipInfo('symbolic-link');entry.create_system=3;entry.external_attr=(0o120777<<16);archive.writestr(entry,'/etc/hosts')
with zipfile.ZipFile(root/'large-entry.zip','w',zipfile.ZIP_DEFLATED) as archive:
    with archive.open('large.txt','w') as output:
        for _ in range(51): output.write(b'A'*(1024*1024))
with wave.open(str(root/'audio.wav'),'wb') as audio:
    audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(8000)
    audio.writeframes(b''.join(struct.pack('<h',round(math.sin(i*2*math.pi*440/8000)*3000)) for i in range(4000)))
for extension,codec in [('mp3','libmp3lame'),('ogg','libvorbis')]:
    subprocess.run(['ffmpeg','-loglevel','error','-y','-i',str(root/'audio.wav'),'-c:a',codec,str(root/f'audio.{extension}')],check=True)
for extension,codec in [('mp4','libx264'),('webm','libvpx-vp9')]:
    subprocess.run(['ffmpeg','-loglevel','error','-y','-loop','1','-i',str(root/'image.png'),'-t','0.5','-vf','format=yuv420p','-c:v',codec,str(root/f'video.{extension}')],check=True)
unsupported = bytearray((root/'audio.wav').read_bytes())
struct.pack_into('<H', unsupported, 20, 0xffff)  # Valid WAV container, unsupported format tag.
(root/'unsupported-codec.wav').write_bytes(unsupported)
print('Generated synthetic preview fixtures in',root)
